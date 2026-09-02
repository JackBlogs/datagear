# -*- coding: utf-8 -*-
"""SDS md → docx：从零创建文档，显式落实格式硬性要求：
- 正文仿宋 12pt（小四）；表格文字仿宋 10.5pt
- 标题 Heading 1（16pt 加粗黑）/ Heading 2（14pt 加粗黑），不带序号；### → 加粗正文段
- 表格：1 磅单实线边框（sz=8，含 insideH/insideV）、表头加粗、无底色
- 封面（28pt/20pt 仿宋加粗居中）+ 文档信息表 + 修订记录表 + TOC 域 + updateFields=true
"""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = sys.argv[1] if len(sys.argv) > 1 else "SDS-软件开发说明书-V1.0.md"
DST = sys.argv[2] if len(sys.argv) > 2 else "SDS-软件开发说明书-V1.0.docx"
FONT = "仿宋"

def set_run_font(run, size_pt, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), FONT)
    rFonts.set(qn('w:hAnsi'), FONT)
    rFonts.set(qn('w:eastAsia'), FONT)

def style_heading(doc, name, size_pt):
    st = doc.styles[name]
    st.font.name = FONT
    st.font.size = Pt(size_pt)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    rPr = st.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        rFonts.set(qn(a), FONT)

def add_para(doc, text, size=12, bold=False, style=None, align=None):
    if style == 'Heading 1':
        size, bold = 16, True
    elif style == 'Heading 2':
        size, bold = 14, True
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    parts = text.split("\t")
    for k, part in enumerate(parts):
        if k > 0:
            r_tab = p.add_run()
            r_tab.add_tab()
            set_run_font(r_tab, size, bold)
        r = p.add_run(part)
        set_run_font(r, size, bold)
    return p

def set_table_borders(table):
    tblPr = table._element.tblPr
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '8')  # 1 磅
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), '000000')
        borders.append(e)
    tblPr.append(borders)
    # 表宽 100%
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:type'), 'pct'); tblW.set(qn('w:w'), '5000')

def add_table(doc, rows):
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = doc.styles['Table Grid']
    set_table_borders(table)
    for ri, row in enumerate(rows):
        tr = table.add_row()
        for j in range(ncols):
            cell = tr.cells[j]
            p = cell.paragraphs[0]
            r = p.add_run(row[j] if j < len(row) else "")
            set_run_font(r, 10.5, bold=(ri == 0))  # 表头加粗
    return table

def add_toc_field(doc):
    p = doc.add_paragraph()
    def mk(tag, **attrs):
        e = OxmlElement(tag)
        for k, v in attrs.items():
            e.set(qn(k), v)
        return e
    r1 = OxmlElement('w:r'); r1.append(mk('w:fldChar', **{'w:fldCharType': 'begin'}))
    r2 = OxmlElement('w:r')
    instr = mk('w:instrText'); instr.set(qn('xml:space'), 'preserve')
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    r2.append(instr)
    r3 = OxmlElement('w:r'); r3.append(mk('w:fldChar', **{'w:fldCharType': 'separate'}))
    r4 = OxmlElement('w:r')
    t = OxmlElement('w:t'); t.text = '（目录将在打开文档时自动更新）'
    r4.append(t)
    r5 = OxmlElement('w:r'); r5.append(mk('w:fldChar', **{'w:fldCharType': 'end'}))
    for r in (r1, r2, r3, r4, r5):
        p._element.append(r)
    return p

def add_page_break(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    br = OxmlElement('w:br'); br.set(qn('w:type'), 'page')
    r._element.append(br)

def is_table_line(s):
    s = s.strip()
    return s.startswith("|") and s.endswith("|") and s.count("|") >= 2

def main():
    lines = open(SRC, encoding="utf-8").read().splitlines()
    doc = Document()
    # A4 页面
    sec = doc.sections[0]
    sec.page_width, sec.page_height = 11906 * 635, 16838 * 635  # twips→EMU
    # 样式
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(12)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        rFonts.set(qn(a), FONT)
    style_heading(doc, 'Heading 1', 16)
    style_heading(doc, 'Heading 2', 14)

    # 封面：前两行为主/副标题
    add_para(doc, lines[0].strip(), size=28, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, lines[1].strip(), size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "", size=12)

    i = 2
    n = len(lines)
    while i < n:
        s = lines[i].strip()
        if not s:
            i += 1; continue
        if is_table_line(s):
            rows = []
            while i < n and is_table_line(lines[i]):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue
        if s.startswith("### "):
            add_para(doc, s[4:].strip(), size=12, bold=True)  # 原文档体系无 H3，用加粗正文段
            i += 1; continue
        if s.startswith("## "):
            add_para(doc, s[3:].strip(), style='Heading 2')
            i += 1; continue
        if s.startswith("# "):
            add_para(doc, s[2:].strip(), style='Heading 1')
            i += 1; continue
        if s == "目录":
            add_para(doc, "目录", size=12, bold=True)
            add_toc_field(doc)
            add_page_break(doc)
            i += 1; continue
        if s == "（此处为 Word 目录域，打开文档后自动更新）":
            i += 1; continue  # 占位行已被 TOC 域取代
        add_para(doc, s)
        i += 1

    # settings.xml: updateFields
    settings = doc.settings.element
    uf = OxmlElement('w:updateFields')
    uf.set(qn('w:val'), 'true')
    settings.insert(0, uf)

    doc.save(DST)
    print(f"saved {DST}: paragraphs={len(doc.paragraphs)}, tables={len(doc.tables)}")

if __name__ == "__main__":
    main()
