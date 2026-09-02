# -*- coding: utf-8 -*-
"""校验 V4.1 docx：1) 元素序列与 prd_v41.md 完全对齐；2) 结构统计；3) 字体抽查；4) 关键词抽检。"""
import re
from docx import Document
from docx.oxml.ns import qn
import align_dryrun as A

# 给 align_dryrun 的 norm 打补丁：剥离 markdown 粗体标记（docx 落地时已去除 **）
_orig_norm = A.norm
def _norm2(s):
    return _orig_norm((s or "").replace("**", ""))
A.norm = _norm2
md_elements = A.md_elements
docx_elements = A.docx_elements
norm = A.norm
key_md = A.key_md
key_dx = A.key_dx

new_md = md_elements("prd_v41.md")
doc = Document("PRD-BI产品需求文档-V4.1-优化版.docx")
dx = docx_elements(doc)
print(f"new_md elems: {len(new_md)} (tbl {sum(1 for e in new_md if e[0]=='tbl')})")
print(f"docx elems: {len(dx)} (p {sum(1 for e in dx if e[0]=='p')}, tbl {sum(1 for e in dx if e[0]=='tbl')})")

# 1. 全量对齐校验（允许 docx 多出空段落 / 文末空段 / h3渲染差异）
fails = []
j = 0
matched = 0
for i, el in enumerate(new_md):
    k = key_md(el)
    found = None
    jj = j
    while jj < len(dx):
        d = dx[jj]
        if d[0] == "p" and not norm(d[2]):
            jj += 1; continue
        dk = key_dx(d)
        if dk == k:
            found = jj; break
        # h3 渲染为加粗正文段：key 前缀不同（md 是 h3 也进 P|），文本应一致
        if el[0] == "h3" and d[0] == "p" and norm(d[2]) == norm(el[1]):
            found = jj; break
        if el[0] == "tbl" and d[0] == "tbl" and len(d[1]) == len(el[1]) and \
           [norm(c) for c in d[1][0]] == [norm(c) for c in el[1][0]]:
            found = jj; break
        jj += 1
    if found is None:
        fails.append((i, el[0], str(el[1])[:70]))
    else:
        matched += 1; j = found + 1
print(f"内容对齐: {matched}/{len(new_md)}, fails: {len(fails)}")
for f_ in fails[:15]:
    print("  FAIL", f_)

# 2. 结构统计
h1 = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
h2 = [p for p in doc.paragraphs if p.style.name == "Heading 2"]
print(f"段落 {len(doc.paragraphs)} / 表格 {len(doc.tables)} / H1 {len(h1)} / H2 {len(h2)}")
empty_tbls = [i for i, t in enumerate(doc.tables) if len(t.rows) == 0]
print("空表格:", len(empty_tbls))
texts = [p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells]
garbled = [t for t in texts if "\ufffd" in t or "\x00" in t]
print("乱码文本:", len(garbled))

# 3. 新内容字体抽查（含新增章节关键词的段落/表格 run 的 eastAsia）
def check_font(keyword):
    hits = []
    for p in doc.paragraphs:
        if keyword in p.text:
            for r in p.runs:
                rPr = r._element.rPr
                ea = rPr.rFonts.get(qn('w:eastAsia')) if rPr is not None and rPr.rFonts is not None else None
                hits.append(("p", p.style.name, ea, r.font.size))
            break
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                if keyword in c.text:
                    for p in c.paragraphs:
                        for r in p.runs:
                            rPr = r._element.rPr
                            ea = rPr.rFonts.get(qn('w:eastAsia')) if rPr is not None and rPr.rFonts is not None else None
                            hits.append(("tbl", "", ea, r.font.size))
                    return hits
    return hits

for kw in ["商业模式与投入测算", "指标字典（新增）", "待采集（Phase 0 埋点上线后回填）", "用户旅程地图（新增）", "数据迁移与升级兼容（新增）"]:
    print(f"字体[{kw}]:", check_font(kw))

# 4. 十五项关键词抽检
all_text = "\n".join(texts)
KW = ["随 Phase 2 交付（与 FR-AI-01 及 Roadmap 一致）", "资源级权限校验", "商业模式与定价", "指标字典",
      "FR-HOME-23", "NFR-AVAIL-08", "标准问题集≥200题", "检索式问答", "达梦/人大金仓双库适配",
      "最小埋点子集", "技术Spike①", "口径统一", "实施工程师", "消费类场景移动可用", "300并发用户/50并发分析任务"]
miss = [k for k in KW if k not in all_text]
print("关键词命中:", len(KW) - len(miss), "/", len(KW), " 缺失:", miss)

# 5. 不应存在的旧内容
for bad in ["V4.0最终整合版），按标准PRD模板", "行级权限透传：同一看板不同用户嵌入看到不同行集", "FR-GOV-17", "数百并发"]:
    print(f"旧内容清除[{bad[:24]}...]:", bad not in all_text)

# 6. updateFields
import zipfile
s = zipfile.ZipFile("PRD-BI产品需求文档-V4.1-优化版.docx").read("word/settings.xml").decode("utf-8")
print("updateFields:", "updateFields" in s)
EOF_MARKER = True
