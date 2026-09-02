# -*- coding: utf-8 -*-
"""将 prd_v41.md 转换为 docx：# →Heading1，## →Heading2，### →Heading3，连续 | ... | 行 → Table Grid 表格，其余为正文段落。"""
import re
from docx import Document
from docx.shared import Pt

SRC = "prd_v41.md"
DST = "PRD-BI产品需求文档-V4.1-优化版.docx"

def is_table_line(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and s.count("|") >= 2

def parse_row(line: str):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells

def clean(text: str) -> str:
    # 去除 markdown 粗体标记
    return text.replace("**", "")

def main():
    with open(SRC, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    # 默认字体
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(clean(line[4:].strip()), level=3)
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(clean(line[3:].strip()), level=2)
            i += 1
            continue
        if line.startswith("# "):
            doc.add_heading(clean(line[2:].strip()), level=1)
            i += 1
            continue
        if is_table_line(line):
            block = []
            while i < n and is_table_line(lines[i]):
                block.append(parse_row(lines[i]))
                i += 1
            ncols = max(len(r) for r in block)
            table = doc.add_table(rows=0, cols=ncols)
            table.style = "Table Grid"
            for r in block:
                row = table.add_row()
                for j in range(ncols):
                    row.cells[j].text = clean(r[j]) if j < len(r) else ""
            continue
        doc.add_paragraph(clean(line.strip()))
        i += 1

    doc.save(DST)
    print(f"saved {DST}: paragraphs={len(doc.paragraphs)}, tables={len(doc.tables)}")

if __name__ == "__main__":
    main()
